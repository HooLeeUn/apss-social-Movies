import csv
import io
from tempfile import SpooledTemporaryFile
from django.http import FileResponse, StreamingHttpResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def safe(value):
    if isinstance(value, str) and value != "-" and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + value
    return value


def export_rows(result, grouped=False):
    yield from result.metadata
    yield []
    if grouped and result.averages:
        yield [value for group in result.header_groups
               for value in [group.label, *([None] * (group.span - 1))]]
        yield [value for group in result.header_groups
               for value in (group.subheaders or (None,))]
    else:
        yield result.headers
    for row in result.rows():
        yield [safe(value) for value in row]


def csv_response(result):
    def chunks():
        yield '\ufeff'
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        for row in export_rows(result):
            writer.writerow([safe(value) for value in row])
            yield buffer.getvalue()
            buffer.seek(0)
            buffer.truncate(0)
    response = StreamingHttpResponse(chunks(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="reccool-report.csv"'
    return response


def xlsx_response(result):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Reporte")
    header_start = len(result.metadata) + 2
    header_end = header_start + int(result.averages)
    bold_font = Font(bold=True)
    sheet.freeze_panes = f"B{header_end + 1}"
    for index, header in enumerate(result.headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(50, max(18, len(header) + 2))
    for index, row in enumerate(export_rows(result, grouped=True), 1):
        cells = []
        for column, value in enumerate(row):
            cell = WriteOnlyCell(sheet, value=safe(value))
            if header_start <= index <= header_end or (index > header_end and row[0] == "TOTAL"):
                cell.font = bold_font
            if index > header_end and result.compares and isinstance(value, (int, float)):
                if column == len(row) - 1:
                    cell.number_format = '0.00"%"'
                elif column == len(row) - 2:
                    cell.number_format = '+0;-0;0'
            cells.append(cell)
        sheet.append(cells)
    methodology = workbook.create_sheet("Metodología")
    for row in result.methodology:
        methodology.append([safe(value) for value in row])
    output = SpooledTemporaryFile(max_size=8*1024*1024, mode="w+b")
    workbook.save(output)
    output.seek(0)
    return FileResponse(output, as_attachment=True, filename="reccool-report.xlsx", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
