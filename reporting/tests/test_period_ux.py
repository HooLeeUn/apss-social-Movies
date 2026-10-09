from datetime import date
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from reporting.forms import ReportForm
from reporting.periods import periods
from reporting.privacy import MESSAGE, variation
from reporting.exports import csv_response, xlsx_response
from reporting.tests.test_reporting import ReportFixtures, dt, filters


class PeriodUXTests(ReportFixtures, TestCase):
    def setUp(self):
        super().setUp()
        self.stamp(self.user, "date_joined", dt(2025, 1))

    def form(self, **kwargs):
        return ReportForm({"report": "users", "months": [1, 2], "years": [2025, 2026], **kwargs})

    def test_cartesian_order_all_monthly_reports(self):
        for report in ["users", "countries", "ages", "genders", "productions", "ratings", "social", "recommended"]:
            form = self.form(report=report, months=[2, 1], years=[2026, 2025])
            self.assertTrue(form.is_valid(), form.errors)
            self.assertEqual([p.label for p in periods(form.cleaned_data)], ["Enero 2025", "Febrero 2025", "Enero 2026", "Febrero 2026"])

    def test_period_limit_and_duplicates(self):
        self.stamp(self.user, "date_joined", dt(2023, 1))
        for years, valid in [([2023, 2024], True), ([2023, 2024, 2025], False)]:
            form = self.form(months=list(range(1, 13)), years=years)
            self.assertEqual(form.is_valid(), valid, form.errors)
        form = self.form(months=[1, 1], years=[2025, 2025])
        self.assertTrue(form.is_valid())
        self.assertEqual(len(periods(form.cleaned_data)), 1)

    def test_future_and_current_month(self):
        with patch("reporting.forms.timezone.localdate", return_value=date(2026, 2, 1)):
            self.assertTrue(self.form(months=[2], years=[2026]).is_valid())
            self.assertFalse(self.form(months=[3], years=[2026]).is_valid())
            self.assertFalse(self.form(months=[1], years=[2027]).is_valid())

    def test_server_modes_and_required_fields(self):
        invalid = [
            {"total": "on", "custom": "on"}, {"total": "on"},
            {"custom": "on", "since": "2026-01-01", "until": "2026-01-31"},
            {"since": "2026-01-01"}, {"months": []}, {"years": []},
            {"months": [13]}, {"years": [1900]},
            {"report": "ratings", "total": "on", "months": [], "years": []},
        ]
        for data in invalid:
            self.assertFalse(self.form(**data).is_valid(), data)
        self.assertTrue(self.form(total="on", months=[], years=[]).is_valid())
        form = self.form(custom="on", months=[], years=[], since="2025-01-15", until="2025-01-15")
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(list(__import__("reporting.services", fromlist=["generate"]).generate(form.cleaned_data).rows()), [["15/01/2025",1,"-","-"],["TOTAL",1,0,0.0]])

    def test_monthly_users_not_accumulated_inactive_and_pending(self):
        from core.models import PendingUserRegistration
        self.user.is_active = False
        self.user.save(update_fields=["is_active"])
        other = get_user_model().objects.create_user(username="feb")
        self.stamp(other, "date_joined", dt(2026, 2))
        PendingUserRegistration.objects.create(username="pending", email="pending@example.com", first_name="A", last_name="B", birth_date=date(2000, 1, 1), password="hashed")
        self.assertEqual(self.rows(total=False, months=[1, 2], years=[2025, 2026]), [["Enero 2025",1,"-","-"],["Febrero 2025",0,-1,-100.0],["Enero 2026",0,0,"No aplica"],["Febrero 2026",1,1,"No aplica"],["TOTAL",2,0,0.0]])
        self.assertEqual(self.rows(), [["Acumulado total",2,"-","-"]])

    def test_user_monthly_privacy_filters_and_derived_cells(self):
        people = self.people(10)
        for person in people:
            self.stamp(person, "date_joined", dt(2025, 1))
        other = get_user_model().objects.create_user(username="minority")
        self.stamp(other, "date_joined", dt(2025, 2))
        for report in ["users", "countries", "ages", "genders"]:
            result = self.result(report, total=False, months=[1, 2], years=[2025], countries=["CO"], genders=["male"], ages=["18_30"])
            self.assertFalse(result.blocked)
            self.assertEqual(list(result.rows())[0][1], 10)
            self.assertTrue(all(cell == MESSAGE for cell in list(result.rows())[1][1:]))
            self.assertTrue(all(cell == MESSAGE for cell in list(result.rows())[-1][1:]))
            self.assertIsNone(result.population)
            csv = b"".join(csv_response(result).streaming_content).decode("utf-8-sig")
            self.assertIn(MESSAGE, csv)
            wb = load_workbook(BytesIO(b"".join(xlsx_response(result).streaming_content)))
            self.assertIn(MESSAGE, [cell for row in wb["Reporte"].values for cell in row])

    def test_clean_exports_methodology_and_population(self):
        result = self.result()
        self.assertFalse(result.show_population)
        self.assertFalse(any("Población" in key for key, _ in result.metadata))
        csv = b"".join(csv_response(result).streaming_content).decode("utf-8-sig")
        self.assertNotIn("Limitación", csv)
        self.assertNotIn("updated_at", csv)
        wb = load_workbook(BytesIO(b"".join(xlsx_response(result).streaming_content)))
        self.assertEqual(wb.sheetnames, ["Reporte", "Metodología"])
        self.assertEqual(list(wb["Metodología"].values), [tuple(row) for row in result.methodology])
        self.assertEqual(list(wb["Reporte"].values)[-1], ("Acumulado total",1,"-","-"))
        self.assertTrue(self.result("ratings").show_population)
        self.assertTrue(self.result(total=False, months=[1, 2], years=[2025]).show_population)

    def test_explicit_variation_headers_and_zero(self):
        for report in ["users", "productions", "comments", "videos", "follows"]:
            result = self.result(report, total=False, months=[1, 2], years=[2025])
            if report == "users":
                self.assertEqual(result.headers[-2:], ["Diferencia", "Variación %"])
            elif report == "productions":
                self.assertEqual(result.headers[-2:], ["Dif. Nº calificaciones", "Var. Nº calificaciones %"])
            else:
                self.assertEqual(len(result.headers), 2)
        self.assertEqual(variation(0, 2), (2, "No aplica"))
        self.assertEqual(variation(1, 1), (0, 0))
        self.assertEqual(variation(2, 0), (-2, -100))

    def test_admin_controls_and_collapsed_methodology(self):
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save(update_fields=["is_staff", "is_superuser"])
        self.client.force_login(self.user)
        response = self.client.get(reverse("admin:reporting_reportaccess_changelist"), {"report": "users", "total": "on"})
        self.assertContains(response, 'name="years"')
        self.assertContains(response, 'name="months"')
        self.assertContains(response, "Combinaciones de géneros más y mejor calificadas")
        self.assertContains(response, "<details><summary>Filtros y semántica")
        self.assertNotContains(response, "Población analizada")
