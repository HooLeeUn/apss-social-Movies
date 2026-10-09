from datetime import date, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import (Comment, CommentReaction, Follow, PendingUserRegistration,
                         VideoComment, VideoCommentReaction)
from reporting.catalog import REPORTS, USER_REPORTS, INTERNAL_REPORTS
from reporting.exports import csv_response, xlsx_response
from reporting.forms import ReportForm
from reporting.periods import midnight, periods
from reporting.privacy import MESSAGE
from reporting.services import generate
from reporting.tests.test_reporting import ReportFixtures, dt, filters


class DailyReportTests(ReportFixtures, TestCase):
    def daily(self, report="users", **kwargs):
        return self.result(report, total=False, custom=True, since=date(2026, 9, 1),
                           until=date(2026, 9, 3), **kwargs)

    def test_inclusive_days_timezone_boundaries_and_long_range(self):
        ps = periods(filters(total=False, custom=True, since=date(2026, 9, 1), until=date(2026, 9, 3)))
        self.assertEqual([p.label for p in ps], ["01/09/2026", "02/09/2026", "03/09/2026"])
        for i, p in enumerate(ps):
            self.assertEqual(p.start, midnight(date(2026, 9, 1) + timedelta(days=i)))
            self.assertEqual(p.end, midnight(date(2026, 9, 2) + timedelta(days=i)))
            self.assertEqual(str(p.start.tzinfo), "America/Bogota")
        self.assertEqual(len(periods(filters(total=False, custom=True, since=date(2026, 9, 1), until=date(2026, 10, 9)))), 39)
        self.assertEqual(len(periods(filters(total=False, custom=True, since=date(2025, 1, 1), until=date(2026, 1, 1)))), 366)
        self.assertEqual(len(periods(filters(total=False, custom=True, since=date(2024, 2, 28), until=date(2024, 3, 1)))), 3)

    def test_all_reports_accept_custom_and_reject_ambiguous_modes(self):
        for report in REPORTS:
            if report in INTERNAL_REPORTS:
                continue  # Internal choices require an explicitly authorized user.
            data = {"report": report, "custom": "on", "since": "2026-09-01", "until": "2026-09-03"}
            with self.subTest(report=report):
                form = ReportForm(data)
                self.assertTrue(form.is_valid(), form.errors)
                self.assertEqual(len(periods(form.cleaned_data)), 3)
                for extra in [{"total": "on"}, {"months": [9]}, {"years": [2026]},
                              {"since": ""}, {"until": ""}, {"until": "2026-08-31"},
                              {"until": "9999-12-31"}]:
                    self.assertFalse(ReportForm({**data, **extra}).is_valid(), extra)
                self.assertEqual(ReportForm({"report": report, "total": "on"}).is_valid(), report in USER_REPORTS)
                self.assertFalse(ReportForm({"report": report, "since": "2026-09-01"}).is_valid())
        self.assertFalse(ReportForm({"report": "social", "custom": "on", "since": "2026-09-01", "until": "2026-09-03", "content_type": "movie"}).is_valid())

    def test_users_daily_inactive_pending_profile_and_midnight(self):
        self.stamp(self.user, "date_joined", datetime(2026, 9, 2, 4, 59, tzinfo=ZoneInfo("UTC")))
        self.user.is_active = False
        self.user.save(update_fields=["is_active"])
        other = get_user_model().objects.create_user(username="day-two")
        self.stamp(other, "date_joined", datetime(2026, 9, 2, 5, tzinfo=ZoneInfo("UTC")))
        outside = get_user_model().objects.create_user(username="outside")
        self.stamp(outside, "date_joined", midnight(date(2026, 9, 4)))
        no_profile = get_user_model().objects.create_user(username="no-profile")
        self.stamp(no_profile, "date_joined", dt(2026, 9, 3))
        no_profile.profile.delete()
        PendingUserRegistration.objects.create(username="pending-daily", email="daily@example.com", first_name="A", last_name="B", birth_date=date(2000, 1, 1), password="hashed")
        self.assertEqual(list(self.daily().rows()), [["01/09/2026",1,"-","-"],["02/09/2026",1,0,0.0],["03/09/2026",0,-1,-100.0],["TOTAL",2,-1,-100.0]])

    def test_daily_demographics_privacy_and_exports(self):
        people = self.people(21, birth=date(2000, 1, 1))
        for i, user in enumerate(people):
            profile = user.profile
            profile.birth_date = date(2000, 1, 1)
            profile.save()
            self.stamp(user, "date_joined", dt(2026, 9, 1 if i < 10 else 2 if i == 10 else 3))
        for report in USER_REPORTS:
            result = self.daily(report, countries=["CO"], ages=["18_30"], genders=["male"])
            self.assertFalse(result.blocked)
            self.assertIsNone(result.population)
            rows = list(result.rows())
            self.assertEqual([row[1] for row in rows], [10, MESSAGE, 10, MESSAGE])
            self.assertEqual(rows[1][-2:], [MESSAGE, MESSAGE])
            self.assertEqual(rows[2][-2:], [MESSAGE, MESSAGE])
            self.assertEqual(rows[-1][-2:], [0, 0.0])
            self.assertEqual(result.metadata[-1][1], MESSAGE)
            csv = b"".join(csv_response(result).streaming_content).decode("utf-8-sig")
            self.assertIn(MESSAGE, csv)
            wb = load_workbook(BytesIO(b"".join(xlsx_response(result).streaming_content)))
            self.assertEqual(list(wb["Reporte"].values)[-1], tuple(rows[-1]))
        for user in people:
            profile = user.profile
            profile.birth_date = date(1995, 9, 2)
            profile.save()
        self.assertEqual([row[1] for row in self.daily(ages=["18_30"]).rows()], [10,MESSAGE,MESSAGE,MESSAGE])

    def test_content_daily_averages_genres_combinations_and_type(self):
        people = self.people(3)
        for user, day, score in zip(people, [1, 2, 2], [8, 6, 10]):
            rating = self.rating(user=user, score=score)
            self.stamp(rating, "updated_at", dt(2026, 9, day))
            self.stamp(rating, "created_at", dt(2026, 8, 1))
        values = [["01/09/2026",1,8.0,1,"-","-"],["02/09/2026",2,8.0,2,1,100.0],["03/09/2026",0,"No aplica",0,-2,-100.0],["TOTAL",3,None,3,-1,-100.0]]
        self.assertEqual(list(self.daily("productions").rows()), values)
        self.assertEqual(list(self.daily("combinations").rows()), values)
        self.assertEqual(list(self.daily("genres").rows()), [["01/09/2026",1,8.0,1,8.0,2,"-","-"],["02/09/2026",2,8.0,2,8.0,4,2,100.0],["03/09/2026",0,"No aplica",0,"No aplica",0,-4,-100.0],["TOTAL",3,None,3,None,6,-2,-100.0]])
        self.assertEqual(list(self.daily("productions", content_type="series").rows()), [["01/09/2026",0,"-","-"],["02/09/2026",0,0,"No aplica"],["03/09/2026",0,0,"No aplica"],["TOTAL",0,0,"No aplica"]])

    def test_recommendations_daily_overlap_and_unique_users(self):
        self.interval(dt(2026, 9, 1), dt(2026, 9, 3))
        self.assertEqual(list(self.result("recommended", total=False, custom=True, since=date(2026, 9, 1), until=date(2026, 9, 4)).rows()), [["01/09/2026",1],["02/09/2026",1],["03/09/2026",1],["04/09/2026",0],["TOTAL",3]])
        self.interval(datetime(2026, 9, 3, 15, tzinfo=ZoneInfo("America/Bogota")), midnight(date(2026, 9, 4)))
        result = self.result("recommended", total=False, custom=True, since=date(2026, 9, 1), until=date(2026, 9, 5))
        # Preserve existing inclusive history end, including exact midnight.
        self.assertEqual(list(result.rows()), [["01/09/2026",1],["02/09/2026",1],["03/09/2026",1],["04/09/2026",1],["05/09/2026",0],["TOTAL",4]])

    def test_direct_daily_metrics_and_consolidated(self):
        self.stamp(self.rating(), "updated_at", dt(2026, 9, 1))
        for visibility in ["public", "mentioned"]:
            comment = Comment.objects.create(author=self.user, movie=self.movie, body="daily", visibility=visibility, is_hidden=True, target_user=self.user if visibility == "mentioned" else None)
            self.stamp(comment, "created_at", dt(2026, 9, 2))
        video = VideoComment.objects.create(user=self.user, movie=self.movie, video="daily.mp4", duration_seconds=1, mime_type="video/mp4", file_size=1, is_hidden=True)
        self.stamp(video, "created_at", dt(2026, 9, 3))
        self.interval(dt(2026, 9, 1), dt(2026, 9, 2))
        rows = list(self.daily("direct").rows())
        self.assertEqual(rows, [["01/09/2026",1,0,0,1,0],["02/09/2026",0,1,0,0,1],["03/09/2026",0,0,1,0,0],["TOTAL",1,1,1,1,1]])
        for i, metric in enumerate(["ratings", "comments", "videos", "added", "removed"], 1):
            self.assertEqual(list(self.daily(metric).rows()), [[row[0],row[i]] for row in rows])

    def test_social_daily_current_state_and_consolidated(self):
        people = self.people(3)
        comment = Comment.objects.create(author=self.user, movie=self.movie, body="daily", is_hidden=True)
        video = VideoComment.objects.create(user=self.user, movie=self.movie, video="daily.mp4", duration_seconds=1, mime_type="video/mp4", file_size=1, is_hidden=True)
        for model, target in [(CommentReaction, {"comment": comment}), (VideoCommentReaction, {"video_comment": video})]:
            for user, reaction_type, day in zip(people[:2], ["like", "dislike"], [1, 2]):
                reaction = model.objects.create(user=user, reaction_type=reaction_type, **target)
                self.stamp(reaction, "updated_at", dt(2026, 9, day))
        follow = Follow.objects.create(follower=self.user, following=people[1])
        self.stamp(follow, "created_at", dt(2026, 9, 3))
        rows = list(self.daily("social").rows())
        self.assertEqual(rows, [["01/09/2026",1,0,1,0,0],["02/09/2026",0,1,0,1,0],["03/09/2026",0,0,0,0,1],["TOTAL",1,1,1,1,1]])
        for i, metric in enumerate(["comment_likes", "comment_dislikes", "video_likes", "video_dislikes", "follows"], 1):
            self.assertEqual(list(self.daily(metric).rows()), [[row[0],row[i]] for row in rows])
        CommentReaction.objects.filter(reaction_type="like").update(reaction_type="dislike", updated_at=dt(2026, 9, 3))
        self.assertEqual(list(self.daily("comment_likes").rows())[0][1], 0)
        Follow.objects.all().delete()
        self.assertEqual(list(self.daily("follows").rows())[2][1], 0)

    def test_activity_content_daily_unique_contributors_privacy(self):
        for i, user in enumerate(self.people(10)):
            self.stamp(self.rating(user=user), "updated_at", dt(2026, 9, 1 if i < 9 else 2))
        for report in ["ratings", "direct", "productions", "genres", "combinations"]:
            result = self.daily(report, countries=["CO"])
            self.assertFalse(result.blocked)
            self.assertIsNone(result.population)
            for row in result.rows():
                self.assertTrue(all(cell in (MESSAGE, "-", None) for cell in row[1:]))

    def test_html_csv_xlsx_share_all_daily_headers_and_metadata(self):
        self.user.is_staff = self.user.is_superuser = True
        self.user.save(update_fields=["is_staff", "is_superuser"])
        self.client.force_login(self.user)
        self.stamp(self.rating(), "updated_at", dt(2026, 9, 1))
        url = reverse("admin:reporting_reportaccess_changelist")
        for report in ["users", "productions", "direct", "social"]:
            params = {"report": report, "custom": "on", "since": "2026-09-01", "until": "2026-09-03"}
            response = self.client.get(url, params)
            self.assertEqual(response.status_code, 200)
            result = response.context_data["result"]
            self.assertEqual(dict(result.metadata)["Periodos"], "01/09/2026; 02/09/2026; 03/09/2026")
            csv = b"".join(self.client.get(url, {**params, "export": "csv"}).streaming_content).decode("utf-8-sig")
            wb = load_workbook(BytesIO(b"".join(self.client.get(url, {**params, "export": "xlsx"}).streaming_content)))
            self.assertEqual(wb.sheetnames, ["Reporte", "Metodología"])
            if not result.averages:
                self.assertIn(tuple(result.headers), list(wb["Reporte"].values))
            for group in result.header_groups:
                self.assertContains(response, group.label)
            for header in result.headers:
                self.assertIn(header, csv)
            self.assertEqual(list(wb["Reporte"].values)[-4:], [tuple(row) for row in result.rows()])
            if result.compares:
                self.assertIn("día consecutivo anterior", str(list(wb["Metodología"].values)))

    def test_39_day_result_all_report_families_and_daily_variation(self):
        people = self.people(8)
        for i, user in enumerate(people):
            self.stamp(user, "date_joined", dt(2026, 9, 1 if i < 3 else 2))
        self.assertEqual(list(self.daily().rows())[1][1:], [5, 2, 66.67])
        self.stamp(self.rating(), "updated_at", dt(2026, 9, 1))
        for report in ["users", "productions", "genres", "recommended", "direct", "social"]:
            result = generate(filters(report, total=False, custom=True, since=date(2026, 9, 1), until=date(2026, 10, 9)))
            self.assertEqual(len(dict(result.metadata)["Periodos"].split("; ")), 39)
            self.assertEqual(list(result.rows())[-2][0], "09/10/2026")
            self.assertEqual(len(list(result.rows())), 40)
            for row in result.rows():
                self.assertEqual(len(row), len(result.headers))
