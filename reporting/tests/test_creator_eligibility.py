import csv
import re
from datetime import date, datetime
from html import unescape
from io import BytesIO, StringIO
from zoneinfo import ZoneInfo

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Comment, CommentReaction, Follow, VideoComment, VideoCommentReaction
from reporting.creator_eligibility import generate_creator_eligibility
from reporting.exports import csv_response, xlsx_response
from reporting.forms import ReportForm
from reporting.periods import periods
from reporting.services import generate
from reporting.tests.test_reporting import ReportFixtures, dt, filters


class CreatorEligibilityTests(ReportFixtures, TestCase):
    def setUp(self):
        super().setUp()
        self.admin = get_user_model().objects.create_superuser(username="internal-admin", email="admin@example.com")
        self.stamp(self.user, "date_joined", dt(2023, 1))
        self.url = reverse("admin:reporting_reportaccess_changelist")

    def params(self, **kwargs):
        return {"report":"creator_eligibility", "custom":"on", "since":"2026-09-01", "until":"2026-09-03", **kwargs}

    def report(self, **kwargs):
        data = filters("creator_eligibility", total=False, custom=True, since=date(2026,9,1),
                       until=date(2026,9,3), months=[9], years=[2026])
        data.update(kwargs)
        return generate(data, user=self.admin)

    def video(self, creator=None, when=None, hidden=False):
        video = VideoComment.objects.create(user=creator or self.user, movie=self.movie,
            video="creator-test.mp4", duration_seconds=1, mime_type="video/mp4", file_size=1, is_hidden=hidden)
        return self.stamp(video, "created_at", when or dt(2026,9,1))

    def react(self, video, actor, kind, when):
        reaction = VideoCommentReaction.objects.create(video_comment=video, user=actor, reaction_type=kind)
        return self.stamp(reaction, "updated_at", when)

    def follower_pool(self, n):
        # Only fixture construction is bulk: no service/model behavior is changed.
        return get_user_model().objects.bulk_create([
            get_user_model()(username=f"fixture-follower-{i}", password="!") for i in range(n)])

    def follow_many(self, followers, target):
        Follow.objects.bulk_create([Follow(follower_id=user.pk, following_id=target.pk) for user in followers])

    def test_superuser_selector_generation_and_both_exports(self):
        self.video()
        self.client.force_login(self.admin)
        response = self.client.get(self.url)
        self.assertContains(response, "Uso interno")
        self.assertContains(response, "Elegibilidad de creadores")
        self.assertContains(response, "Mínimo de seguidores")
        response = self.client.get(self.url, self.params())
        self.assertContains(response, self.user.username)
        self.assertContains(response, "Usuarios evaluados")
        self.assertContains(response, "Usuarios que cumplen umbral")
        self.assertNotContains(response, "Población analizada")
        self.assertEqual(response.context_data["result"].summary[0], ("Usuarios evaluados",1))
        for fmt in ["csv", "xlsx"]:
            exported = self.client.get(self.url, self.params(export=fmt))
            self.assertEqual(exported.status_code,200)
            self.assertTrue(b"".join(exported.streaming_content))

    def test_normal_analyst_view_and_export_permissions_do_not_grant_internal_access(self):
        analyst = get_user_model().objects.create_user(username="analyst-internal",is_staff=True)
        for code in ["can_view_reports", "can_export_reports"]:
            analyst.user_permissions.add(Permission.objects.get(codename=code,content_type__app_label="reporting"))
        self.client.force_login(analyst)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code,200)
        self.assertNotContains(response,"Elegibilidad de creadores")
        self.assertNotContains(response,"Uso interno")
        self.assertNotContains(response,"min_followers")
        for extra in [{}, {"export":"csv"}, {"export":"xlsx"}]:
            self.assertEqual(self.client.get(self.url,self.params(**extra)).status_code,403)
        self.assertEqual(self.client.get(self.url,{"report":"users","total":"on","export":"csv"}).status_code,200)
        with self.assertNumQueries(0), self.assertRaises(PermissionDenied):
            generate(filters("creator_eligibility",total=False),user=analyst)

    def test_staff_with_only_one_permission_cannot_see_or_export_internal_report(self):
        for code in ["can_view_reports", "can_export_reports"]:
            staff = get_user_model().objects.create_user(username=f"only-{code}",is_staff=True)
            staff.user_permissions.add(Permission.objects.get(codename=code,content_type__app_label="reporting"))
            self.client.force_login(staff)
            response = self.client.get(self.url)
            self.assertNotIn(b"Elegibilidad de creadores",response.content)
            for extra in [{}, {"export":"csv"}, {"export":"xlsx"}]:
                self.assertEqual(self.client.get(self.url,self.params(**extra)).status_code,403)

    def test_nonstaff_inactive_and_anonymous_denied(self):
        for user in [self.user, self.admin]:
            user.is_staff = False
            user.save(update_fields=["is_staff"])
            self.client.force_login(user)
            self.assertEqual(self.client.get(self.url,self.params()).status_code,302)
            with self.assertRaises(PermissionDenied):
                generate(filters("creator_eligibility",total=False),user=user)
        self.admin.is_staff = True
        self.admin.is_active = False
        self.admin.save(update_fields=["is_staff","is_active"])
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(self.url,self.params(export="csv")).status_code,302)
        self.client.logout()
        self.assertEqual(self.client.get(self.url,self.params()).status_code,302)
        with self.assertNumQueries(0), self.assertRaises(PermissionDenied):
            generate(filters("creator_eligibility",total=False))
        with self.assertRaises(PermissionDenied):
            generate_creator_eligibility(filters("creator_eligibility",total=False),None)

    def test_threshold_positive_configurable_default_and_server_modes(self):
        for value in [None,1000,5000,10000,25000,100000]:
            params = self.params()
            if value is not None:
                params["min_followers"] = value
            form = ReportForm(params,user=self.admin)
            self.assertTrue(form.is_valid(),form.errors)
            self.assertEqual(form.cleaned_data["min_followers"],value or 10000)
        for value in [0,-1,"1.5","invalid"]:
            self.assertFalse(ReportForm(self.params(min_followers=value),user=self.admin).is_valid())
        for extra in [{"total":"on"},{"months":[9]},{"years":[2026]}, {"since":""},
                      {"countries":["CO"]},{"ages":["18_30"]},{"genders":["male"]},{"content_type":"movie"}]:
            self.assertFalse(ReportForm(self.params(**extra),user=self.admin).is_valid(),extra)
        self.assertFalse(ReportForm(self.params()).is_valid())
        self.assertEqual(ReportForm(user=self.admin).fields["min_followers"].initial,10000)
        monthly = ReportForm({"report":"creator_eligibility","months":[10,9],"years":[2026]},user=self.admin)
        self.assertTrue(monthly.is_valid(),monthly.errors)
        self.assertEqual([p.label for p in periods(monthly.cleaned_data)],["Septiembre 2026","Octubre 2026"])
        too_many = ReportForm({"report":"creator_eligibility","months":list(range(1,13)),"years":[2023,2024,2025]},user=self.admin)
        self.assertFalse(too_many.is_valid())
        future = ReportForm({"report":"creator_eligibility","months":[12],"years":[2026]},user=self.admin)
        self.assertFalse(future.is_valid())

    def test_large_followers_fixtures_classification_and_configurable_threshold(self):
        creator_b = get_user_model().objects.create_user(username="creator-b")
        creator_c = get_user_model().objects.create_user(username="creator-c")
        followers = self.follower_pool(15000)
        self.follow_many(followers[:12000],self.user)
        self.follow_many(followers[:8000],creator_b)
        self.follow_many(followers,creator_c)
        video_a = self.video()
        self.video(creator_b)
        self.react(video_a,creator_b,"like",dt(2026,9,1))
        rows = list(self.report().rows())
        for creator, count, eligible in [(self.user,12000,"Sí"),(creator_b,8000,"No"),(creator_c,15000,"Sí")]:
            own = [row for row in rows if row[2] == creator.pk]
            self.assertEqual(len(own),3)
            self.assertTrue(all(row[3] == count and row[-1] == eligible for row in own))
        self.assertEqual(rows[0][2],self.user.pk)  # Engagement precedes followers among eligible creators.
        self.assertEqual(self.report().summary[:2],[("Usuarios evaluados",3),("Usuarios que cumplen umbral",2)])
        high = list(self.report(min_followers=25000).rows())
        self.assertEqual({row[2] for row in high},{self.user.pk,creator_b.pk})
        self.assertTrue(all(row[-1] == "No" for row in high))

    def test_followers_9999_10000_10001_boundary(self):
        creators = self.people(3)
        followers = self.follower_pool(10001)
        for creator, count in zip(creators,[9999,10000,10001]):
            self.follow_many(followers[:count],creator)
            self.video(creator)
        rows = list(self.report(until=date(2026,9,1)).rows())
        self.assertEqual({row[3]:row[-1] for row in rows},{9999:"No",10000:"Sí",10001:"Sí"})

    def test_follow_direction_current_snapshot_and_threshold_only_creator(self):
        other = get_user_model().objects.create_user(username="follow-target")
        self.video()
        self.stamp(Follow.objects.create(follower=self.user,following=other),"created_at",dt(2020,1,1))
        result = self.report(min_followers=1)
        rows = list(result.rows())
        self.assertEqual([row[3] for row in rows if row[2] == self.user.pk],[0,0,0])
        self.assertTrue(all(row[3] == 1 and row[-1] == "Sí" for row in rows if row[2] == other.pk))
        Follow.objects.all().delete()
        self.assertTrue(all(row[3] == 1 for row in result.rows() if row[2] == other.pk))
        self.assertNotIn(other.pk,{row[2] for row in self.report(min_followers=1).rows()})

    def test_received_likes_dislikes_are_attributed_to_creator_not_actors(self):
        actor_b = get_user_model().objects.create_user(username="reactor-b")
        actor_c = get_user_model().objects.create_user(username="reactor-c")
        video = self.video()
        self.video(actor_b,dt(2025,1))  # Include actor as a creator, with no received engagement.
        self.react(video,actor_b,"like",dt(2026,9,1))
        self.react(video,actor_c,"dislike",dt(2026,9,1))
        rows = list(self.report().rows())
        self.assertEqual(next(row for row in rows if row[2] == self.user.pk)[5:8],[1,1,2])
        self.assertTrue(all(row[5:8] == [0,0,0] for row in rows if row[2] == actor_b.pk))
        self.assertNotIn(actor_c.pk,{row[2] for row in rows})
        self.assertTrue(all(row[-1] == "No" for row in rows))

    def test_daily_bogota_midnight_state_change_and_deletion(self):
        actor = get_user_model().objects.create_user(username="timestamp-reactor")
        video = self.video(when=datetime(2026,9,2,4,59,tzinfo=ZoneInfo("UTC")),hidden=True)
        reaction = self.react(video,actor,"like",datetime(2026,9,2,5,0,tzinfo=ZoneInfo("UTC")))
        rows = list(self.report().rows())
        self.assertEqual([row[0] for row in rows],["01/09/2026","02/09/2026","03/09/2026"])
        self.assertEqual([row[4:8] for row in rows],[[1,0,0,0],[0,1,0,1],[0,0,0,0]])
        VideoCommentReaction.objects.filter(pk=reaction.pk).update(reaction_type="dislike",updated_at=dt(2026,9,3))
        self.assertEqual([row[4:8] for row in self.report().rows()],[[1,0,0,0],[0,0,0,0],[0,0,1,1]])
        reaction.delete()
        self.assertTrue(all(row[7] == 0 for row in self.report().rows()))

    def test_monthly_videos_and_received_reactions_separate(self):
        actor = get_user_model().objects.create_user(username="monthly-reactor")
        september = self.video(when=dt(2026,9,15))
        october = self.video(when=dt(2026,10,1))
        self.react(september,actor,"like",dt(2026,9,15))
        self.react(october,actor,"dislike",dt(2026,10,1))
        result = self.report(custom=False,since=None,until=None,months=[10,9],years=[2026])
        self.assertEqual(list(result.rows()),[["Septiembre 2026",self.user.username,self.user.pk,0,1,1,0,1,"No"],
                                           ["Octubre 2026",self.user.username,self.user.pk,0,1,0,1,1,"No"]])

    def test_include_old_video_inactive_staff_but_exclude_superuser_and_unrelated_accounts(self):
        self.user.is_staff = True
        self.user.is_active = False
        self.user.username = "technical"  # No arbitrary username-based exclusions.
        self.user.save()
        self.video(when=dt(2025,1))
        self.video(self.admin)
        unrelated = get_user_model().objects.create_user(username="unrelated")
        rows = list(self.report().rows())
        self.assertEqual({row[2] for row in rows},{self.user.pk})
        self.assertTrue(all(row[4:8] == [0,0,0,0] for row in rows))
        self.assertNotIn(unrelated.pk,{row[2] for row in rows})

    def test_public_comment_reactions_never_enter_metrics(self):
        self.video()
        actor = get_user_model().objects.create_user(username="comment-reactor")
        comment = Comment.objects.create(author=self.user,movie=self.movie,body="public")
        reaction = CommentReaction.objects.create(user=actor,comment=comment,reaction_type="like")
        self.stamp(reaction,"updated_at",dt(2026,9,1))
        self.assertTrue(all(row[5:8] == [0,0,0] for row in self.report().rows()))

    def test_order_eligible_then_engagement_followers_and_username(self):
        creators = self.people(4)
        for creator, name in zip(creators,["zeta","alpha","high-engagement","qualified"]):
            creator.username = name
            creator.save()
            self.video(creator)
        actors = self.follower_pool(3)
        self.follow_many(actors[:2],creators[0])
        self.follow_many(actors[:2],creators[1])
        self.follow_many(actors,creators[3])
        high_video = VideoComment.objects.get(user=creators[2])
        for actor in actors:
            self.react(high_video,actor,"like",dt(2026,9,1))
        rows = list(self.report(min_followers=2,until=date(2026,9,1)).rows())
        self.assertEqual([row[1] for row in rows],["qualified","alpha","zeta","high-engagement"])

    def test_html_csv_xlsx_same_metrics_no_personal_fields_or_totals(self):
        self.user.email = "creator-private@example.com"
        self.user.first_name = "PrivateLegalName"
        self.user.save()
        self.video()
        self.client.force_login(self.admin)
        response = self.client.get(self.url,self.params())
        result = response.context_data["result"]
        rows = list(result.rows())
        body = response.content.decode().split("<tbody>",1)[1].split("</tbody>",1)[0]
        html_rows = [[unescape(value).strip() for value in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>",row)]
                     for row in re.findall(r"<tr[^>]*>(.*?)</tr>",body,flags=re.S)]
        self.assertEqual(html_rows,[[str(value) for value in row] for row in rows])
        csv_data = b"".join(csv_response(result).streaming_content).decode("utf-8-sig")
        csv_rows = list(csv.reader(StringIO(csv_data)))
        self.assertEqual(csv_rows[csv_rows.index(result.headers)+1:],[[str(value) for value in row] for row in rows])
        wb = load_workbook(BytesIO(b"".join(xlsx_response(result).streaming_content)))
        self.assertEqual(list(wb["Reporte"].values)[-3:],[tuple(row) for row in rows])
        self.assertEqual(wb.sheetnames,["Reporte","Metodología"])
        self.assertIn("unfollows históricos",str(list(wb["Metodología"].values)))
        self.assertIn("derecho a compensación",str(list(wb["Metodología"].values)))
        for text in [response.content.decode(),csv_data,str(list(wb["Reporte"].values))]:
            self.assertNotIn(self.user.email,text)
            self.assertNotIn(self.user.first_name,text)
            self.assertNotIn("TOTAL",text)
        self.assertFalse(result.compares)
        self.assertFalse(result.show_population)
        self.assertEqual(len(rows),3)  # No TOTAL and no privacy suppression at n=1.

    def test_detail_pagination_exports_all_rows_and_no_n_plus_one(self):
        creators = self.people(30)
        for creator in creators:
            self.video(creator)
        with self.assertNumQueries(7):
            result = self.report()
            full = list(result.rows())
        self.assertEqual(len(full),90)
        with self.assertNumQueries(2):
            page = list(result.rows(50,10))
        self.assertEqual(page,full[50:60])
        self.client.force_login(self.admin)
        response = self.client.get(self.url,self.params(page=2))
        self.assertEqual(response.context_data["rows"],full[50:])
        self.assertEqual(response.context_data["previous_page"],1)
        export = self.client.get(self.url,self.params(page=2,export="csv"))
        csv_rows = list(csv.reader(StringIO(b"".join(export.streaming_content).decode("utf-8-sig"))))
        self.assertEqual(len(csv_rows[csv_rows.index(result.headers)+1:]),90)

    def test_zero_candidates_no_data_and_service_rejects_invalid_parameters(self):
        result = self.report()
        self.assertEqual(result.summary[0],("Usuarios evaluados",0))
        self.assertEqual(list(result.rows()),[])
        for kwargs in [{"min_followers":0},{"min_followers":"10"},{"total":True},{"countries":["CO"]}]:
            with self.assertRaises(ValueError):
                self.report(**kwargs)
