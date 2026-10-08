from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from zoneinfo import ZoneInfo
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient
from openpyxl import load_workbook
from core.models import (Movie, MovieRating, Profile, PendingUserRegistration, Comment, VideoComment, CommentReaction, VideoCommentReaction, Follow, MovieRecommendationItem, MovieRecommendationHistory, UserGenrePreference, UserTypePreference, UserDirectorPreference, UserTasteProfile)
from core.recommendation_history import backfill_history
from reporting.services import generate
from reporting.demographics import age_on
from reporting.privacy import MESSAGE, variation
from reporting.forms import ReportForm
from reporting.exports import csv_response, xlsx_response

TZ = ZoneInfo("America/Bogota")
def dt(year, month, day=15):
    return datetime(year, month, day, 12, tzinfo=TZ)


def filters(report="users", **kwargs):
    data = {"report": report, "countries": [], "ages": [], "genders": [], "content_type": "", "total": True, "custom": False, "month": 1, "year": 2026, "since": None, "until": None, "months": ["2026-01"], **kwargs}
    if "years" not in kwargs:
        if report == "users" or report in {"countries", "ages", "genders"}:
            data["months"] = [data["month"]]
            data["years"] = [data["year"]]
        else:
            pairs = [value.split("-") for value in data["months"]]
            data["months"] = sorted({int(m) for y, m in pairs})
            data["years"] = sorted({int(y) for y, m in pairs})
    return data


class ReportFixtures:
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="report-user")
        self.user.profile.birth_date = date(1995, 6, 1)
        self.user.profile.gender_identity = "male"
        self.user.profile.save()
        self.movie = Movie.objects.create(title_english="Dune", type="movie", genre="Drama, Action, Drama", director="Test director")
        self.api = APIClient()
        self.api.force_authenticate(self.user)

    def result(self, report="users", **kwargs):
        return generate(filters(report, **kwargs))

    def rows(self, report="users", **kwargs):
        return list(self.result(report, **kwargs).rows())

    def stamp(self, obj, field, value):
        type(obj).objects.filter(pk=obj.pk).update(**{field: value})
        obj.refresh_from_db()
        return obj

    def rating(self, user=None, movie=None, score=8, month=1):
        r = MovieRating.objects.create(user=user or self.user, movie=movie or self.movie, score=score)
        return self.stamp(r, "updated_at", dt(2026, month))

    def people(self, n, country="CO", gender="male", birth=date(1995, 6, 1)):
        people = [self.user]
        for i in range(n-1):
            u = get_user_model().objects.create_user(username=f"person-{i}")
            p = u.profile; p.streaming_country=country; p.gender_identity=gender; p.birth_date=birth; p.save()
            people.append(u)
        return people

    def interval(self, start, end=None, user=None):
        return MovieRecommendationHistory.objects.create(user=user or self.user, movie=self.movie, started_at=start, ended_at=end)

class ReportTests(ReportFixtures, TestCase):
    def test_pending_does_not_count(self):
        PendingUserRegistration.objects.create(username="pending", email="a@b.com", first_name="A", last_name="B", birth_date=date(2000,1,1), password="hashed")
        self.assertEqual(self.rows(), [["Usuarios registrados", 1]])

    def test_confirmed_counts_and_date_joined(self):
        pending = PendingUserRegistration.objects.create(username="confirmed", email="a@b.com", first_name="A", last_name="B", birth_date=date(2000,1,1), password="hashed", terms_accepted_at=dt(2025,1))
        response = self.api.get(reverse("register-confirm-email", args=[pending.token]))
        self.assertEqual(response.status_code, 302)
        user = get_user_model().objects.get(username="confirmed")
        self.assertEqual(user.date_joined.date(), timezone.now().date())
        self.assertEqual(user.profile.terms_accepted_at, dt(2025,1))
        self.assertEqual(self.result().population, 2)

    def test_inactive_registered_counts(self):
        self.user.is_active=False; self.user.save()
        self.assertEqual(self.result().population, 1)

    def test_user_without_profile_excluded(self):
        self.user.profile.delete()
        self.assertEqual(self.result().population, 0)

    def test_country_filters_single_multiple_and_combined(self):
        qs = __import__("reporting.queries", fromlist=["users"]).users
        period = __import__("reporting.periods", fromlist=["Period"]).Period("all",None,None)
        from reporting.periods import month_period
        self.stamp(self.user, "date_joined", dt(2026,1))
        for countries, count in [(["CO"],1),(["CO","VE"],1),(["VE"],0)]:
            self.assertEqual(qs(filters(countries=countries),period).count(),count)
        self.assertEqual(qs(filters(countries=["CO","VE"], genders=["male"], ages=["18_30"]),month_period("2026-01")).count(),1)
        self.assertEqual(qs(filters(genders=["female"]),period).count(),0)

    def test_all_age_bands_and_historical_age(self):
        from reporting.queries import users
        from reporting.periods import month_period
        self.stamp(self.user,"date_joined",dt(2026,1))
        for years, key in [(15,"13_17"),(30,"18_30"),(35,"31_40"),(45,"41_50"),(55,"51_plus")]:
            p=self.user.profile; p.birth_date=date(2026-years,1,1); p.save()
            self.assertEqual(users(filters(ages=[key]),month_period("2026-01")).count(),1)
        p=self.user.profile; p.birth_date=date(1995,6,1); p.save()
        self.assertEqual(users(filters(ages=["18_30"]),month_period("2026-01")).count(),1)
        self.assertEqual(users(filters(ages=["31_40"]),month_period("2026-01")).count(),0)

    def test_birthday_and_leap_year_python_sql_match(self):
        from reporting.queries import users
        from reporting.periods import Period
        for birth, reference, expected in [(date(1995,6,1),date(2026,5,31),30),(date(1995,6,1),date(2026,6,1),31),(date(2000,2,29),date(2025,2,28),24),(date(2000,2,29),date(2025,3,1),25),(date(2000,2,29),date(2024,2,29),24)]:
            self.assertEqual(age_on(birth,reference),expected)
            p=self.user.profile; p.birth_date=birth;p.save()
            self.stamp(self.user,"date_joined",datetime.combine(reference,datetime.min.time(),TZ))
            self.assertEqual(users(filters(),Period("all",dt(1900,1),dt(2090,1))).get().report_age,expected)

    def test_user_accumulated_month_custom_periods(self):
        self.stamp(self.user,"date_joined",dt(2026,1,31))
        self.assertEqual(self.result().population,1)
        self.assertEqual(self.result(total=False,month=1).population,1)
        self.assertEqual(self.result(total=False,month=2).population,0)
        self.assertEqual(self.result(total=False,custom=True,since=date(2026,1,31),until=date(2026,1,31)).population,1)

    def test_user_group_reports(self):
        self.assertEqual(self.rows("countries"),[])
        self.assertEqual(self.result("countries").blocked,True)
        self.people(10)
        self.assertEqual(self.rows("countries"),[["Colombia",10]])
        self.assertEqual(self.rows("genders"),[["Hombre",10]])
        self.assertEqual(self.rows("ages")[0][1],10)

    def test_rating_updated_at_average_and_content_type(self):
        self.rating(score=8,month=2)
        other=get_user_model().objects.create_user(username="other")
        self.rating(user=other,score=6,month=2)
        self.assertEqual(self.rows("productions",months=["2026-01"]),[])
        self.assertEqual(self.rows("productions",months=["2026-02"]),[["Dune","Película",2,7.0]])
        self.assertEqual(self.rows("productions",months=["2026-02"],content_type="series"),[])
        series=Movie.objects.create(title_english="Series",type="series")
        self.rating(movie=series,month=2)
        self.assertEqual(self.rows("productions",months=["2026-02"],content_type="series")[0][:3],["Series","Serie",1])

    def test_genres_individual_and_exact_combination(self):
        self.rating()
        self.assertEqual(self.rows("genres"),[["Action",1,8.0],["Drama",1,8.0]])
        self.assertEqual(self.rows("combinations"),[["Action|Drama",1,8.0]])

    def test_content_multiple_months_zero_and_variation(self):
        self.rating()
        u=get_user_model().objects.create_user(username="u2");self.rating(user=u,month=2)
        u=get_user_model().objects.create_user(username="u3");self.rating(user=u,month=2)
        row=self.rows("productions",months=["2026-01","2026-02","2026-03"])[0]
        self.assertEqual(row,["Dune","Película",1,8.0,2,8.0,1,100.0,0,"No aplica",-2,-100.0])
        self.assertEqual(variation(0,3),(3,"No aplica"))

    def test_public_comments_only_hidden_count(self):
        for public, hidden in [(True,False),(True,True),(False,False)]:
            c=Comment.objects.create(author=self.user,movie=self.movie,body="test",visibility="public" if public else "mentioned",is_hidden=hidden,target_user=self.user if not public else None)
            self.stamp(c,"created_at",dt(2026,1))
        self.assertEqual(self.rows("comments"),[["Comentarios públicos realizados",2]])

    def test_hidden_video_counts(self):
        v=VideoComment.objects.create(user=self.user,movie=self.movie,video="test.mp4",duration_seconds=1,mime_type="video/mp4",file_size=1,is_hidden=True)
        self.stamp(v,"created_at",dt(2026,1))
        self.assertEqual(self.rows("videos"),[["Video reacciones publicadas",1]])

    def test_reactions_and_follows_actor_and_dates(self):
        other=get_user_model().objects.create_user(username="target")
        comment=Comment.objects.create(author=other,movie=self.movie,body="public",is_hidden=True)
        video=VideoComment.objects.create(user=other,movie=self.movie,video="a.mp4",duration_seconds=1,mime_type="video/mp4",file_size=1,is_hidden=True)
        for model,target in [(CommentReaction,{"comment":comment}),(VideoCommentReaction,{"video_comment":video})]:
            r=model.objects.create(user=self.user,reaction_type="like",**target);self.stamp(r,"updated_at",dt(2026,1))
        f=Follow.objects.create(follower=self.user,following=other);self.stamp(f,"created_at",dt(2026,1))
        rows=self.rows("social")
        self.assertEqual([r[1] for r in rows],[1,0,1,0,1])
        for model in [CommentReaction,VideoCommentReaction]:
            model.objects.update(reaction_type="dislike",updated_at=dt(2026,2))
        self.assertEqual([r[1] for r in self.rows("social",months=["2026-02"])],[0,1,0,1,0])
        comment.visibility="mentioned";comment.save()
        self.assertEqual(self.rows("comment_dislikes",months=["2026-02"])[0][1],0)

    def test_direct_multiple_months_and_activity_updated_at(self):
        r=self.rating(month=2)
        self.stamp(r,"created_at",dt(2026,1))
        rows=self.rows("direct",months=["2026-01","2026-02"])
        self.assertEqual(rows[0],["Calificaciones realizadas",0,1,1,"No aplica"])
        self.assertEqual(len(rows),5)

    def test_privacy_9_blocked_10_allowed_and_exports(self):
        people=self.people(9)
        for u in people:self.rating(user=u)
        r=self.result("ratings",countries=["CO"])
        self.assertTrue(r.blocked);self.assertIsNone(r.population);self.assertEqual(list(r.rows()),[])
        u=get_user_model().objects.create_user(username="tenth");self.rating(user=u)
        r=self.result("ratings",countries=["CO"])
        self.assertFalse(r.blocked);self.assertEqual(list(r.rows())[0][1],10)

    def test_privacy_per_cell_and_no_derived_disclosure(self):
        people=self.people(10)
        for u in people:self.rating(user=u)
        for u in people[:9]:
            c=Comment.objects.create(author=u,movie=self.movie,body="test");self.stamp(c,"created_at",dt(2026,2))
        r=self.result("direct",countries=["CO"],months=["2026-01","2026-02"])
        self.assertFalse(r.blocked)
        row=list(r.rows())[0]
        self.assertEqual(row,["Calificaciones realizadas",10,MESSAGE,MESSAGE,MESSAGE])
        self.assertEqual(list(r.rows())[1],["Comentarios públicos realizados",MESSAGE,MESSAGE,MESSAGE,MESSAGE])
        data=b"".join(csv_response(r).streaming_content).decode("utf-8-sig")
        self.assertIn(MESSAGE,data)
        wb=load_workbook(BytesIO(b"".join(xlsx_response(r).streaming_content)))
        self.assertIn(MESSAGE,[v for row in wb.active.values for v in row])

    def test_privacy_content_each_production_and_genre(self):
        people=self.people(10)
        for u in people:self.rating(user=u)
        minority=Movie.objects.create(title_english="Hidden title",genre="Secret",type="movie")
        self.rating(movie=minority)
        r=self.result("productions",countries=["CO"])
        rows=list(r.rows());self.assertEqual(rows[0][2],10)
        self.assertEqual(rows[1],[MESSAGE,MESSAGE,MESSAGE,MESSAGE])
        self.assertEqual(self.rows("genres",countries=["CO"])[-1],[MESSAGE,MESSAGE,MESSAGE])

    def test_csv_bom_and_real_xlsx_and_formula_safety(self):
        self.movie.title_english="=HYPERLINK(1)";self.movie.save();self.rating()
        r=self.result("productions")
        data=b"".join(csv_response(r).streaming_content)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"));self.assertIn(b"'=HYPERLINK",data)
        x=b"".join(xlsx_response(r).streaming_content);self.assertTrue(x.startswith(b"PK"))
        wb=load_workbook(BytesIO(x));last=list(wb.active.rows)[-1][0]
        self.assertEqual(last.data_type,"s");self.assertTrue(last.value.startswith("'="))

    def test_form_validation_and_irrelevant_fields(self):
        form=ReportForm({"report":"users","total":"on","content_type":"movie"})
        self.assertTrue(form.is_valid());self.assertEqual(form.cleaned_data["content_type"],"")
        self.assertFalse(ReportForm({"report":"ratings"}).is_valid())
        self.assertFalse(ReportForm({"report":"users","custom":"on","since":"2026-02-01","until":"2026-01-01"}).is_valid())
        self.assertFalse(ReportForm({"report":"users","total":"on","countries":["invalid"]}).is_valid())

    def test_report_pagination(self):
        for i in range(55):
            m=Movie.objects.create(title_english=f"Film-{i}",type="movie");self.rating(movie=m)
        r=self.result("productions")
        self.assertEqual(len(list(r.rows(0,50))),50)
        self.assertEqual(len(list(r.rows(50,50))),5)


class RatingIdempotenceTests(ReportFixtures, TestCase):
    def test_same_score_does_not_save_or_change_preferences(self):
        url=reverse("movie-rating",args=[self.movie.pk])
        response=self.api.put(url,{"score":8},format="json")
        self.assertEqual(response.json(),{"movie":self.movie.pk,"my_rating":8,"created":True})
        rating=MovieRating.objects.get();self.stamp(rating,"updated_at",dt(2026,1))
        before={m.__name__:list(m.objects.values()) for m in [UserGenrePreference,UserTypePreference,UserDirectorPreference,UserTasteProfile]}
        with patch("core.signals.update_user_preferences_for_movie_rating") as sync:
            response=self.api.put(url,{"score":8},format="json");sync.assert_not_called()
        self.assertEqual(response.json(),{"movie":self.movie.pk,"my_rating":8,"created":False})
        rating.refresh_from_db();self.assertEqual(rating.updated_at,dt(2026,1))
        self.assertEqual(MovieRating.objects.count(),1)
        for m in [UserGenrePreference,UserTypePreference,UserDirectorPreference,UserTasteProfile]:
            self.assertEqual(before[m.__name__],list(m.objects.values()))
        self.assertEqual(self.rows("productions")[0][2],1)
        response=self.api.put(url,{"score":9},format="json")
        self.assertEqual(response.json(),{"movie":self.movie.pk,"my_rating":9,"created":False})
        rating.refresh_from_db();self.assertGreater(rating.updated_at,dt(2026,1))
        self.assertEqual(MovieRating.objects.count(),1)
        for m in [UserGenrePreference,UserTypePreference,UserDirectorPreference]:
            for pref in m.objects.all():
                self.assertEqual((pref.count_8,pref.count_9,pref.ratings_count,pref.score),(0,1,1,Decimal("9.00")))
        self.assertEqual(UserTasteProfile.objects.get().ratings_count,1)


class RecommendationTests(ReportFixtures, TestCase):
    def test_add_remove_readd_idempotent_preserves_intervals(self):
        url=reverse("movie-recommendation-toggle",args=[self.movie.pk])
        self.assertTrue(self.api.post(url).json()["created"])
        self.assertEqual(MovieRecommendationItem.objects.count(),1)
        self.assertEqual(MovieRecommendationHistory.objects.count(),1)
        self.assertFalse(self.api.post(url).json()["created"])
        self.assertEqual(MovieRecommendationHistory.objects.count(),1)
        self.assertTrue(self.api.delete(url).json()["deleted"])
        self.assertEqual(MovieRecommendationItem.objects.count(),0)
        self.assertIsNotNone(MovieRecommendationHistory.objects.get().ended_at)
        self.assertFalse(self.api.delete(url).json()["deleted"])
        self.api.post(url)
        self.assertEqual(MovieRecommendationHistory.objects.count(),2)
        self.assertEqual(MovieRecommendationHistory.objects.filter(ended_at=None).count(),1)

    def test_overlap_january_february_not_march(self):
        self.interval(dt(2026,1,10),dt(2026,2,12))
        self.assertEqual(self.rows("recommended",months=["2026-01","2026-02","2026-03"]),[["Dune","Película",1,1,0,0.0,0,-1,-100.0]])

    def test_multiple_intervals_count_once_same_month(self):
        self.interval(dt(2026,1,3),dt(2026,1,7));self.interval(dt(2026,1,20),dt(2026,1,25))
        self.assertEqual(self.rows("recommended"),[["Dune","Película",1]])
        self.assertEqual(self.rows("added")[0][1],2)
        self.assertEqual(self.rows("removed")[0][1],2)
        self.assertEqual(self.rows("recommended",months=["2026-02"]),[])

    def test_backfill_preserves_timestamp_idempotent_partial(self):
        item=MovieRecommendationItem.objects.create(user=self.user,movie=self.movie)
        self.stamp(item,"created_at",dt(2025,1))
        MovieRecommendationHistory.objects.all().delete()
        self.assertEqual(backfill_history(),1)
        self.assertEqual(MovieRecommendationHistory.objects.get().started_at,dt(2025,1))
        self.assertIsNone(MovieRecommendationHistory.objects.get().ended_at)
        self.assertEqual(backfill_history(),0)
        self.assertEqual(MovieRecommendationHistory.objects.count(),1)

    def test_overlap_exact_month_boundary(self):
        from reporting.periods import month_period
        feb=month_period("2026-02")
        self.interval(dt(2026,1),feb.start)
        self.assertEqual(self.rows("recommended",months=["2026-02"])[0][2],1)

    def test_user_deletion_does_not_recreate_history(self):
        MovieRecommendationItem.objects.create(user=self.user,movie=self.movie)
        self.user.delete()
        self.assertEqual(MovieRecommendationHistory.objects.count(),0)


class PermissionTests(TestCase):
    def setUp(self):
        self.staff=get_user_model().objects.create_user(username="analyst",password="test",is_staff=True)
        self.url=reverse("admin:reporting_reportaccess_changelist")
        self.client.force_login(self.staff)

    def grant(self,code):
        self.staff.user_permissions.add(Permission.objects.get(codename=code,content_type__app_label="reporting"))

    def test_staff_without_permission_direct_url_denied(self):
        self.assertEqual(self.client.get(self.url).status_code,403)
        self.assertNotContains(self.client.get(reverse("admin:index")),"Generar reportes")

    def test_view_permission_and_only_report_menu(self):
        self.grant("can_view_reports")
        self.assertEqual(self.client.get(self.url).status_code,200)
        index=self.client.get(reverse("admin:index"));self.assertContains(index,"Generar reportes");self.assertNotContains(index,"/admin/core/")
        self.assertContains(self.client.get(self.url,{"report":"users","total":"on"}),"Usuarios registrados")

    def test_export_requires_both_permissions_direct_url(self):
        self.grant("can_view_reports")
        for fmt in ["csv","xlsx"]:
            self.assertEqual(self.client.get(self.url,{"report":"users","total":"on","export":fmt}).status_code,403)
        self.grant("can_export_reports")
        for fmt in ["csv","xlsx"]:
            response=self.client.get(self.url,{"report":"users","total":"on","export":fmt});self.assertEqual(response.status_code,200);b"".join(response.streaming_content)

    def test_export_only_cannot_view(self):
        self.grant("can_export_reports");self.assertEqual(self.client.get(self.url).status_code,403)

    def test_privacy_blocked_export_direct_urls(self):
        self.grant("can_view_reports");self.grant("can_export_reports")
        for fmt in ["csv","xlsx"]:
            self.assertEqual(self.client.get(self.url,{"report":"users","total":"on","countries":["CO"],"export":fmt}).status_code,403)

    def test_nonstaff_inactive_and_anonymous_denied(self):
        self.grant("can_view_reports")
        self.staff.is_staff=False;self.staff.save();self.assertEqual(self.client.get(self.url).status_code,302)
        self.staff.is_staff=True;self.staff.is_active=False;self.staff.save();self.assertEqual(self.client.get(self.url).status_code,302)
        self.client.logout();self.assertEqual(self.client.get(self.url).status_code,302)

    def test_superuser_full_access(self):
        self.staff.is_superuser=True;self.staff.save()
        self.assertEqual(self.client.get(self.url).status_code,200)

    def test_group_setup_idempotent(self):
        from django.core.management import call_command
        from django.contrib.auth.models import Group
        call_command("setup_report_analysts");call_command("setup_report_analysts")
        self.assertEqual(Group.objects.get(name="Analistas de reportes").permissions.count(),2)

class AdditionalReportTests(ReportFixtures, TestCase):
    def test_multiple_countries_combined_demographics_report(self):
        people=self.people(10)
        for i,u in enumerate(people):
            p=u.profile;p.streaming_country="CO" if i<5 else "VE";p.save()
            self.stamp(u,"date_joined",dt(2026,1))
        result=self.result(countries=["CO","VE"],genders=["male"],ages=["18_30"],total=False)
        self.assertEqual(result.population,10)
        self.assertEqual(list(result.rows()),[["Usuarios registrados",10]])
        self.assertTrue(self.result(countries=["CO"],total=False).blocked)

    def test_population_not_exposed_with_suppressed_demographic_group(self):
        people=self.people(11)
        p=people[-1].profile;p.streaming_country="VE";p.save()
        r=self.result("countries")
        self.assertIsNone(r.population)
        self.assertEqual(list(r.rows()),[["Colombia",10],["Venezuela",MESSAGE]])
        self.assertEqual(r.metadata[-1][1],MESSAGE)

    def test_age_unknown_and_outside_groups(self):
        people=self.people(10)
        Profile.objects.update(birth_date=None)
        self.assertEqual(self.rows("ages"),[["Sin dato",10]])
        Profile.objects.update(birth_date=date(2020,1,1))
        self.assertEqual(self.rows("ages"),[["Fuera de rangos definidos",10]])

    def test_timezone_month_boundaries(self):
        rating=self.rating()
        # 04:59 UTC February 1 still belongs to January in Bogota.
        self.stamp(rating,"updated_at",datetime(2026,2,1,4,59,tzinfo=ZoneInfo("UTC")))
        self.assertEqual(self.rows("ratings")[0][1],1)
        self.assertEqual(self.rows("ratings",months=["2026-02"])[0][1],0)
        self.stamp(rating,"updated_at",datetime(2026,2,1,5,0,tzinfo=ZoneInfo("UTC")))
        self.assertEqual(self.rows("ratings")[0][1],0)
        self.assertEqual(self.rows("ratings",months=["2026-02"])[0][1],1)

    def test_event_age_filter_not_current_age(self):
        people=self.people(10)
        for u in people:self.rating(user=u)
        r=self.result("ratings",ages=["18_30"])
        self.assertFalse(r.blocked);self.assertEqual(list(r.rows())[0][1],10)
        self.assertTrue(self.result("ratings",ages=["31_40"]).blocked)

    def test_recommended_age_on_overlap_reference(self):
        from reporting.queries import recommendations
        from reporting.periods import month_period
        self.interval(dt(2024,1),None)
        self.assertEqual(recommendations(filters(ages=["18_30"]),month_period("2026-01")).count(),1)
        self.assertEqual(recommendations(filters(ages=["18_30"]),month_period("2026-07")).count(),0)
        self.assertEqual(recommendations(filters(ages=["31_40"]),month_period("2026-07")).count(),1)

    def test_remove_before_backfill_recovers_real_start(self):
        item=MovieRecommendationItem.objects.create(user=self.user,movie=self.movie)
        self.stamp(item,"created_at",dt(2025,1))
        MovieRecommendationHistory.objects.all().delete()
        self.api.delete(reverse("movie-recommendation-toggle",args=[self.movie.pk]))
        history=MovieRecommendationHistory.objects.get()
        self.assertEqual(history.started_at,dt(2025,1));self.assertIsNotNone(history.ended_at)
        self.assertEqual(backfill_history(),0)

    def test_movie_deletion_cascades_history(self):
        MovieRecommendationItem.objects.create(user=self.user,movie=self.movie)
        self.movie.delete()
        self.assertEqual(MovieRecommendationHistory.objects.count(),0)

    def test_recommendation_add_rolls_back_if_history_fails(self):
        from core.recommendation_history import set_recommendation
        with patch("core.signals.MovieRecommendationHistory.objects.get_or_create",side_effect=RuntimeError("test failure")):
            with self.assertRaises(RuntimeError):set_recommendation(self.user,self.movie.pk,True)
        self.assertEqual(MovieRecommendationItem.objects.count(),0)

    def test_recommendation_remove_rolls_back_if_history_fails(self):
        from core.recommendation_history import set_recommendation
        set_recommendation(self.user,self.movie.pk,True)
        with patch("core.signals.MovieRecommendationHistory.objects.get_or_create",side_effect=RuntimeError("test failure")):
            with self.assertRaises(RuntimeError):set_recommendation(self.user,self.movie.pk,False)
        self.assertEqual(MovieRecommendationItem.objects.count(),1)
        self.assertIsNone(MovieRecommendationHistory.objects.get().ended_at)

    def test_hidden_social_demographic_actor(self):
        from reporting.queries import source
        from reporting.periods import month_period
        other=get_user_model().objects.create_user(username="creator")
        p=other.profile;p.streaming_country="VE";p.save()
        c=Comment.objects.create(author=other,movie=self.movie,body="test",is_hidden=True)
        r=CommentReaction.objects.create(user=self.user,comment=c,reaction_type="like");self.stamp(r,"updated_at",dt(2026,1))
        qs,_=source("comment_likes",filters(countries=["CO"]),month_period("2026-01"));self.assertEqual(qs.count(),1)
        qs,_=source("comment_likes",filters(countries=["VE"]),month_period("2026-01"));self.assertEqual(qs.count(),0)

    def test_form_rejects_custom_nonmonthly_and_excess_comparison(self):
        self.assertFalse(ReportForm({"report":"ratings","custom":"on","since":"2026-01-01","until":"2026-01-31"}).is_valid())
        self.assertFalse(ReportForm({"report":"social","months":["2026-01"],"content_type":"movie"}).is_valid())
        form=ReportForm({"report":"users","total":"on","months":["99"]})
        self.assertFalse(form.is_valid())


from django.test import TransactionTestCase
from django.db import close_old_connections
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

class ConcurrentWriteTests(TransactionTestCase):
    def setUp(self):
        self.user=get_user_model().objects.create_user(username="parallel")
        self.movie=Movie.objects.create(title_english="Concurrent",genre="Action",type="movie")

    def parallel(self, operations):
        gate=Barrier(len(operations))
        def execute(operation):
            close_old_connections()
            try:
                client=APIClient();client.force_authenticate(get_user_model().objects.get(pk=self.user.pk))
                gate.wait(timeout=10)
                return operation(client)
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=len(operations)) as pool:
            return list(pool.map(execute,operations))

    def test_parallel_rating_first_writes_one_row_and_preferences(self):
        url=reverse("movie-rating",args=[self.movie.pk])
        results=self.parallel([lambda c:c.put(url,{"score":8},format="json").json()]*2)
        self.assertEqual(sorted(r["created"] for r in results),[False,True])
        self.assertEqual(MovieRating.objects.count(),1)
        self.assertEqual(UserTasteProfile.objects.get().ratings_count,1)
        self.assertEqual(UserGenrePreference.objects.get().count_8,1)

    def test_parallel_recommendation_adds_and_add_remove_consistency(self):
        url=reverse("movie-recommendation-toggle",args=[self.movie.pk])
        results=self.parallel([lambda c:c.post(url).json()]*2)
        self.assertEqual(sorted(r["created"] for r in results),[False,True])
        self.assertEqual(MovieRecommendationHistory.objects.count(),1)
        self.parallel([lambda c:c.delete(url).json(),lambda c:c.post(url).json()])
        self.assertEqual(MovieRecommendationHistory.objects.filter(ended_at=None).count(),MovieRecommendationItem.objects.count())
        self.assertEqual(MovieRecommendationHistory.objects.filter(ended_at__isnull=False).count(),1)

class HistoryConstraintTests(ReportFixtures, TestCase):
    def test_only_one_open_interval_and_valid_end(self):
        from django.db import IntegrityError, transaction
        self.interval(dt(2026,1))
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.interval(dt(2026,2))
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.interval(dt(2026,3),dt(2026,2))
        self.assertEqual(MovieRecommendationHistory.objects.count(),1)

    def test_backfill_does_not_duplicate_closed_intervals_or_change_current_state(self):
        self.interval(dt(2025,1),dt(2025,2))
        item=MovieRecommendationItem.objects.create(user=self.user,movie=self.movie)
        original=MovieRecommendationItem.objects.values().get(pk=item.pk)
        self.assertEqual(backfill_history(),0)
        self.assertEqual(MovieRecommendationHistory.objects.count(),2)
        self.assertEqual(MovieRecommendationItem.objects.values().get(pk=item.pk),original)
