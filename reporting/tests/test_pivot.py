import csv
import re
from datetime import date, datetime
from html import unescape
from io import BytesIO, StringIO
from zoneinfo import ZoneInfo

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import (Comment, CommentReaction, Follow, Movie, VideoComment,
                         VideoCommentReaction)
from reporting.catalog import DIRECT, SOCIAL
from reporting.exports import csv_response, xlsx_response
from reporting.privacy import MESSAGE, protected_sum
from reporting.tests.test_reporting import ReportFixtures, dt


class PivotTests(ReportFixtures, TestCase):
    def daily(self, report="users", days=3, **kwargs):
        return self.result(report, total=False, custom=True, since=date(2026, 9, 1),
                           until=date(2026, 9, days), **kwargs)

    def registrations(self, counts, monthly=False):
        people = iter(self.people(sum(counts)))
        for period, count in enumerate(counts, 1):
            for _ in range(count):
                self.stamp(next(people), "date_joined", dt(2026, period) if monthly else dt(2026, 9, period))

    def export_data(self, result):
        csv_rows = list(csv.reader(StringIO(b"".join(csv_response(result).streaming_content).decode("utf-8-sig"))))
        start = csv_rows.index(result.headers) + 1
        wb = load_workbook(BytesIO(b"".join(xlsx_response(result).streaming_content)))
        expected = list(result.rows())
        self.assertEqual(csv_rows[start:], [["" if value is None else str(value) for value in row] for row in expected])
        self.assertEqual(list(wb["Reporte"].values)[-len(expected):], [tuple(row) for row in expected])
        self.assertEqual(wb.sheetnames, ["Reporte", "Metodología"])
        return wb

    def admin(self):
        self.user.is_staff = self.user.is_superuser = True
        self.user.save(update_fields=["is_staff", "is_superuser"])
        self.client.force_login(self.user)
        return reverse("admin:reporting_reportaccess_changelist")

    def test_registered_users_example_and_first_last_total(self):
        self.registrations([6, 5, 1, 2, 8])
        result = self.daily(days=5)
        self.assertEqual(result.headers, ["Fecha", "Registros", "Diferencia", "Variación %"])
        self.assertEqual(list(result.rows()), [
            ["01/09/2026", 6, "-", "-"], ["02/09/2026", 5, -1, -16.67],
            ["03/09/2026", 1, -4, -80.0], ["04/09/2026", 2, 1, 100.0],
            ["05/09/2026", 8, 6, 300.0], ["TOTAL", 22, 2, 33.33],
        ])
        wb = self.export_data(result)
        self.assertEqual(list(wb["Reporte"].rows)[-1][-1].number_format, '0.00"%"')

    def test_zero_first_and_consecutive_empty_periods(self):
        self.registrations([0, 0, 5])
        result = self.daily()
        self.assertEqual(list(result.rows()), [["01/09/2026",0,"-","-"],
            ["02/09/2026",0,0,"No aplica"], ["03/09/2026",5,5,"No aplica"],
            ["TOTAL",5,5,"No aplica"]])
        self.export_data(result)

    def test_monthly_users_same_orientation_and_first_last_rule(self):
        self.registrations([6, 5, 1, 2, 8], monthly=True)
        result = self.result(total=False, months=[1,2,3,4,5], years=[2026])
        self.assertEqual(result.headers[0], "Periodo")
        self.assertEqual([row[0] for row in result.rows()], ["Enero 2026", "Febrero 2026", "Marzo 2026", "Abril 2026", "Mayo 2026", "TOTAL"])
        self.assertEqual(list(result.rows())[-1], ["TOTAL",22,2,33.33])
        self.export_data(result)
        single = self.result(total=False, months=[1], years=[2026])
        self.assertEqual(list(single.rows()), [["Enero 2026",6,"-","-"],["TOTAL",6,0,0.0]])

    def test_segment_columns_logical_order_and_additive_totals(self):
        for i, user in enumerate(self.people(40)):
            first = i % 20 < 10
            profile = user.profile
            profile.streaming_country = "CO" if first else "VE"
            profile.gender_identity = "male" if first else "female"
            profile.birth_date = date(2000, 1, 1) if first else date(1990, 1, 1)
            profile.save()
            self.stamp(user, "date_joined", dt(2026, 9, 1 if i < 20 else 2))
        for report in ["countries", "ages", "genders"]:
            result = self.daily(report, days=2)
            self.assertEqual(len(result.columns), 2)
            self.assertEqual(result.headers[-3:], ["TOTAL", "Diferencia", "Variación %"])
            self.assertEqual(list(result.rows()), [["01/09/2026",10,10,20,"-","-"],
                ["02/09/2026",10,10,20,0,0.0], ["TOTAL",20,20,40,0,0.0]])
            self.export_data(result)
        self.assertEqual([column.label for column in self.daily("ages", days=2).columns], ["18–30", "31–40"])

    def test_hidden_segment_suppresses_row_column_and_grand_totals(self):
        for i, user in enumerate(self.people(22)):
            profile = user.profile
            profile.streaming_country = "VE" if i % 11 == 10 else "CO"
            profile.save()
            self.stamp(user, "date_joined", dt(2026, 9, 1 if i < 11 else 2))
        result = self.daily("countries", days=2)
        self.assertIsNone(result.population)
        self.assertEqual(list(result.rows()), [["01/09/2026",10,MESSAGE,MESSAGE,"-","-"],
            ["02/09/2026",10,MESSAGE,MESSAGE,MESSAGE,MESSAGE],
            ["TOTAL",20,MESSAGE,MESSAGE,MESSAGE,MESSAGE]])
        self.export_data(result)

    def test_hidden_first_period_suppresses_endpoint_comparisons(self):
        self.registrations([9, 10])
        result = self.daily(days=2, countries=["CO"])
        self.assertFalse(result.blocked)
        self.assertEqual(list(result.rows()), [["01/09/2026",MESSAGE,"-","-"],
            ["02/09/2026",10,MESSAGE,MESSAGE], ["TOTAL",MESSAGE,MESSAGE,MESSAGE]])
        self.export_data(result)

    def test_ratings_pivot_counts_averages_and_count_only_totals(self):
        self.movie.genre = "Action"
        self.movie.save()
        drama = Movie.objects.create(title_english="Drama film", type="movie", genre="Drama")
        events = [(1,self.movie,8), (1,self.movie,6), (1,drama,10),
                  (2,self.movie,4), (2,drama,6), (3,self.movie,7),
                  (3,self.movie,8), (3,self.movie,9), (3,drama,9)]
        for user, (day, movie, score) in zip(self.people(len(events)), events):
            self.stamp(self.rating(user=user, movie=movie, score=score), "updated_at", dt(2026, 9, day))
        expected = [["01/09/2026",2,7.0,1,10.0,3,"-","-"],
            ["02/09/2026",1,4.0,1,6.0,2,-1,-33.33],
            ["03/09/2026",3,8.0,1,9.0,4,2,100.0], ["TOTAL",6,None,3,None,9,1,33.33]]
        for report in ["productions", "genres", "combinations"]:
            result = self.daily(report)
            self.assertEqual(list(result.rows()), expected)
            self.assertEqual(result.header_groups[1].subheaders, ("No calif.", "Promedio Calif."))
            for row in expected[:-1]:
                self.assertEqual(row[1] + row[3], row[5])
            wb = self.export_data(result)
            table = list(wb["Reporte"].values)
            offset = len(result.metadata) + 1
            self.assertEqual(table[offset][1], result.columns[0].label)
            self.assertEqual(table[offset + 1][1:5], ("No calif.", "Promedio Calif.", "No calif.", "Promedio Calif."))
            self.assertIsNone(table[-1][2])
            self.assertIsNone(table[-1][4])

    def test_genre_overlap_total_is_contributions_not_distinct_ratings(self):
        self.stamp(self.rating(score=9), "updated_at", dt(2026, 9, 1))
        result = self.daily("genres", days=1)
        self.assertEqual(list(result.rows()), [["01/09/2026",1,9.0,1,9.0,2,"-","-"], ["TOTAL",1,None,1,None,2,0,0.0]])
        self.assertIn("No equivale", str(result.methodology))

    def test_content_privacy_masks_averages_titles_totals_and_comparisons(self):
        for user in self.people(10):
            self.stamp(self.rating(user=user), "updated_at", dt(2026, 9, 1))
        secret = Movie.objects.create(title_english="Private production", genre="Private genre", type="movie")
        self.stamp(self.rating(movie=secret), "updated_at", dt(2026, 9, 1))
        for report in ["productions", "genres", "combinations"]:
            result = self.daily(report, days=2, countries=["CO"])
            self.assertNotIn("Private", str(result.headers))
            rows = list(result.rows())
            self.assertEqual(rows[0][-3:], [MESSAGE,"-","-"])
            self.assertEqual(rows[1][-3:], [MESSAGE,MESSAGE,MESSAGE])
            self.assertEqual(rows[-1][-3:], [MESSAGE,MESSAGE,MESSAGE])
            self.assertTrue(all(value in (None, MESSAGE) for value in rows[-1][1:-3]))
            self.export_data(result)

    def test_recommended_example_sum_and_no_comparisons(self):
        people = self.people(11)
        for day, count in [(1,8), (2,4), (3,11)]:
            for user in people[:count]:
                self.interval(dt(2026, 9, day), datetime(2026,9,day,23,tzinfo=ZoneInfo("America/Bogota")), user=user)
        result = self.daily("recommended")
        self.assertEqual(result.headers, ["Fecha", "Dune (Película)"])
        self.assertEqual(list(result.rows()), [["01/09/2026",8],["02/09/2026",4],["03/09/2026",11],["TOTAL",23]])
        self.assertEqual(result.population, 11)
        self.export_data(result)
        protected = self.daily("recommended", countries=["CO"])
        self.assertEqual(list(protected.rows()), [["01/09/2026",MESSAGE],["02/09/2026",MESSAGE],["03/09/2026",11],["TOTAL",MESSAGE]])
        self.export_data(protected)

    def test_all_activity_reports_html_csv_xlsx_same_values_without_comparisons(self):
        self.stamp(self.rating(), "updated_at", dt(2026, 9, 1))
        comment = Comment.objects.create(author=self.user, movie=self.movie, body="test", is_hidden=True)
        self.stamp(comment, "created_at", dt(2026, 9, 2))
        video = VideoComment.objects.create(user=self.user, movie=self.movie, video="pivot.mp4", duration_seconds=1, mime_type="video/mp4", file_size=1, is_hidden=True)
        self.stamp(video, "created_at", dt(2026, 9, 3))
        self.interval(dt(2026, 9, 1), dt(2026, 9, 2))
        reaction = CommentReaction.objects.create(user=self.user, comment=comment, reaction_type="like")
        self.stamp(reaction, "updated_at", dt(2026, 9, 1))
        reaction = VideoCommentReaction.objects.create(user=self.user, video_comment=video, reaction_type="dislike")
        self.stamp(reaction, "updated_at", dt(2026, 9, 2))
        other = self.people(2)[1]
        self.stamp(Follow.objects.create(follower=self.user, following=other), "created_at", dt(2026, 9, 3))
        url = self.admin()
        for report in ["direct", *DIRECT, "social", *SOCIAL]:
            result = self.daily(report)
            self.assertFalse(result.compares)
            self.assertFalse(any("Dif" in header or "Var" in header for header in result.headers))
            rows = list(result.rows())
            self.assertEqual(rows[-1], ["TOTAL", *[sum(row[i] for row in rows[:-1]) for i in range(1,len(result.headers))]])
            response = self.client.get(url, {"report":report,"custom":"on","since":"2026-09-01","until":"2026-09-03"})
            self.assertEqual(response.status_code,200)
            body = response.content.decode().split("<tbody>",1)[1].split("</tbody>",1)[0]
            html_rows = [[unescape(re.sub(r"<[^>]+>", "", value)).strip()
                          for value in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row)]
                         for row in re.findall(r"<tr[^>]*>(.*?)</tr>", body, flags=re.S)]
            self.assertEqual(html_rows, [[str(value) for value in row] for row in rows])
            self.export_data(result)

    def test_dynamic_pages_keep_all_periods_and_full_report_totals(self):
        for i in range(55):
            movie = Movie.objects.create(title_english=f"Film {i}", type="movie")
            self.stamp(self.rating(movie=movie), "updated_at", dt(2026, 9, 1))
        url = self.admin()
        params = {"report":"productions","custom":"on","since":"2026-09-01","until":"2026-09-03"}
        first = self.client.get(url, params)
        second = self.client.get(url, {**params,"page":2})
        full = self.daily("productions")
        for response, count in [(first,50),(second,5)]:
            result = response.context_data["result"]
            self.assertEqual(len(result.columns),count)
            self.assertEqual([row[0] for row in result.rows()], ["01/09/2026","02/09/2026","03/09/2026","TOTAL"])
            self.assertEqual([row[-3:] for row in result.rows()], [row[-3:] for row in full.rows()])
            self.assertContains(response, "TOTAL y sus comparaciones abarcan todas las entidades")
        self.assertEqual(first.context_data["next_page"],2)
        self.assertEqual(second.context_data["previous_page"],1)
        exported = self.client.get(url,{**params,"page":2,"export":"csv"})
        csv_rows = list(csv.reader(StringIO(b"".join(exported.streaming_content).decode("utf-8-sig"))))
        self.assertIn(full.headers,csv_rows)
        self.assertEqual(self.client.get(url,{**params,"page":999}).context_data["page"],2)

    def test_off_page_protected_cell_suppresses_global_totals(self):
        people = self.people(10)
        for i in range(51):
            movie = Movie.objects.create(title_english=f"Public {i}", type="movie")
            for user in people:
                self.rating(user=user,movie=movie)
        hidden = Movie.objects.create(title_english="Hidden off page", type="movie")
        self.rating(movie=hidden)
        result = self.result("productions",countries=["CO"])
        first = result.for_columns(0,50)
        self.assertNotIn(MESSAGE, [column.label for column in first.columns])
        self.assertEqual(list(first.rows())[0][-3],MESSAGE)
        self.assertEqual(list(first.rows())[-1][-3:], [MESSAGE,MESSAGE,MESSAGE])
        self.export_data(result)

    def test_equal_titles_disambiguated_and_order_stable_across_formats(self):
        duplicate = Movie.objects.create(title_english="Dune",type="movie",genre="Comedy")
        self.rating(movie=duplicate)
        self.rating()
        result = self.result("productions")
        self.assertEqual(len(set(column.label for column in result.columns)),2)
        self.assertTrue(all("[#" in column.label for column in result.columns))
        self.assertEqual(result.headers, self.result("productions").headers)
        self.export_data(result)

    def test_grouped_html_and_flat_csv_headers_and_blank_total_averages(self):
        self.stamp(self.rating(), "updated_at", dt(2026,9,1))
        url = self.admin()
        response = self.client.get(url,{"report":"genres","custom":"on","since":"2026-09-01","until":"2026-09-03"})
        self.assertContains(response,'scope="colgroup" colspan="2">Action')
        self.assertContains(response,'scope="col">No calif.')
        self.assertContains(response,'class="report-total"')
        self.assertNotContains(response," · Total")
        result = self.daily("genres")
        self.assertIn("Action - No calif.",result.headers)
        self.assertIn("Action - Promedio Calif.",result.headers)
        self.export_data(result)

    def test_protected_sum_never_uses_partial_values(self):
        self.assertEqual(protected_sum([3,0,2]),5)
        self.assertIsNone(protected_sum([10,None,10]))
        self.assertEqual(protected_sum([]),0)
